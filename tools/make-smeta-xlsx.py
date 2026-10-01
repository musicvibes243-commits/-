#!/usr/bin/env python3
"""Собирает docs/smeta.xlsx — бланк сметы для Excel.

Зачем отдельно от smeta.html: страницу удобно заполнять с телефона
на объекте, а файл Excel — присылать заказчику, править на ноутбуке
и хранить у себя папкой. Обе бумаги про одно и то же, просто разные руки.

Суммы в файле — живые формулы, а не числа: поправили цену, итог
пересчитался сам. Поэтому при правке нельзя стирать столбец «Сумма».

Запуск:  python3 tools/make-smeta-xlsx.py
"""

import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from xlsx_formuly import dopisat_rezultaty

OUT = pathlib.Path("docs/smeta.xlsx")

# Цвета взяты из бланка smeta.html, чтобы бумаги выглядели одной семьёй.
INK = "FF23303A"
MUTED = "FF7A8792"
ACCENT = "FFC97A09"
LINE = "FFD8D2C7"
LINE_2 = "FFECE7DE"

# Позиции из разговора 1 октября 2026. Это заготовка: лишнее удаляется,
# своё дописывается. Пустых строк ниже оставлено с запасом.
POZICII = [
    ("Котёл", 1, "шт", 49950),
    ("Монтаж", 1, "усл.", 49000),
    ("Расходный материал для котла", 1, "компл.", 13000),
    ("Отверстия", 2, "шт", 9500),
    ("Монтаж воздуховодов", 1, "усл.", 16000),
    ("Материал", 1, "компл.", 18000),
    ("Доставка материала", 1, "усл.", 7000),
]
ZAPAS = 8          # сколько пустых строк дописать под свои позиции
DENGI = '#,##0.00" ₽"'

thin = Side(style="thin", color=LINE_2)
ramka = Border(left=thin, right=thin, top=thin, bottom=thin)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"

    # Готовые ответы формул — иначе быстрый просмотр на телефоне и в почте
    # показывает нули вместо сумм. Подробности в tools/xlsx_formuly.py.
    otvety = {}

    shirina = {"A": 5, "B": 46, "C": 9, "D": 10, "E": 14, "F": 16}
    for col, w in shirina.items():
        ws.column_dimensions[col].width = w

    def sliyanie(row, text, **kw):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(**kw) if kw else Font(color=INK)
        return c

    # ---------- Шапка ----------
    sliyanie(1, "СтройИнвест", bold=True, size=18, color=INK)
    sliyanie(2, "Инженерные системы дома и участка: септики, водоснабжение, "
                "отопление, электрика, дренаж", size=9, color=MUTED)
    sliyanie(3, "ООО «ИнженерИнвест» · ИНН 2100032923 · ОГРН 1262100002878",
             size=9, color=MUTED)
    sliyanie(4, "+7 (915) 346-97-28 · stroyinvest-mo.ru", size=9, color=MUTED)
    ws.row_dimensions[1].height = 24

    # ---------- Поля заказчика ----------
    polya = [
        (6, "СМЕТА №", ""),
        (7, "Дата", ""),
        (8, "Заказчик", ""),
        (9, "Телефон", ""),
        (10, "Адрес объекта", ""),
        (11, "Объект", ""),
    ]
    for row, podpis, znach in polya:
        c = ws.cell(row=row, column=1, value=podpis)
        c.font = Font(size=9, color=MUTED, bold=(row == 6))
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        v = ws.cell(row=row, column=2, value=znach)
        v.font = Font(size=11, color=INK, bold=(row == 6))
        v.border = Border(bottom=Side(style="dotted", color=LINE))

    # ---------- Таблица ----------
    SHAPKA = 13
    zagolovki = ["№", "Наименование работ, материалов", "Кол-во", "Ед.", "Цена, ₽", "Сумма, ₽"]
    for i, t in enumerate(zagolovki, start=1):
        c = ws.cell(row=SHAPKA, column=i, value=t)
        c.font = Font(bold=True, size=9, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = ramka
    ws.row_dimensions[SHAPKA].height = 26

    pervaya = SHAPKA + 1
    vsego = len(POZICII) + ZAPAS
    for i in range(vsego):
        r = pervaya + i
        poz = POZICII[i] if i < len(POZICII) else ("", None, "", None)
        ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=poz[0] or None)
        ws.cell(row=r, column=3, value=poz[1]).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4, value=poz[2] or None).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5, value=poz[3]).number_format = DENGI
        # Живая формула: поправили цену — сумма и итог пересчитались сами.
        s = ws.cell(row=r, column=6, value=f"=IF(OR(C{r}=\"\",E{r}=\"\"),\"\",C{r}*E{r})")
        s.number_format = DENGI
        otvety[f"F{r}"] = poz[1] * poz[3] if i < len(POZICII) else ""
        for col in range(1, 7):
            cell = ws.cell(row=r, column=col)
            cell.border = ramka
            if cell.font.size is None:
                cell.font = Font(size=11, color=INK)
        ws.row_dimensions[r].height = 19

    poslednya = pervaya + vsego - 1

    # ---------- Итого ----------
    ITOG = poslednya + 1
    ws.merge_cells(start_row=ITOG, start_column=1, end_row=ITOG, end_column=5)
    c = ws.cell(row=ITOG, column=1, value="ИТОГО")
    c.font = Font(bold=True, size=13, color=INK)
    c.alignment = Alignment(horizontal="right")
    otvety[f"F{ITOG}"] = sum(k * c for _, k, _, c in POZICII)
    it = ws.cell(row=ITOG, column=6, value=f"=SUM(F{pervaya}:F{poslednya})")
    it.font = Font(bold=True, size=13, color=INK)
    it.number_format = DENGI
    it.border = Border(top=Side(style="medium", color=INK))
    c.border = Border(top=Side(style="medium", color=INK))
    ws.row_dimensions[ITOG].height = 24

    # ---------- Низ ----------
    def blok(row, zagolovok, vysota=1):
        c = ws.cell(row=row, column=1, value=zagolovok)
        c.font = Font(bold=True, size=9, color=ACCENT)
        ws.merge_cells(start_row=row + 1, start_column=1,
                       end_row=row + vysota, end_column=6)
        v = ws.cell(row=row + 1, column=1)
        v.alignment = Alignment(wrap_text=True, vertical="top")
        v.font = Font(size=10, color=INK)
        v.border = ramka
        return row + vysota + 1

    r = ITOG + 2
    r = blok(r, "СРОКИ", 1)
    r += 1
    r = blok(r, "ЧТО ВХОДИТ В СТОИМОСТЬ", 2)
    r += 1
    r = blok(r, "ЧТО В СТОИМОСТЬ НЕ ВХОДИТ", 2)
    r += 1

    sliyanie(r, "Смета действительна 14 дней с даты составления. Окончательная "
                "стоимость фиксируется договором. Объём работ может измениться, "
                "если на объекте откроются обстоятельства, которые не были видны "
                "при осмотре, — изменения согласуются с заказчиком до начала работ.",
             size=9, color=MUTED)
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 42

    podpis = r + 2
    for col, text in ((1, "Исполнитель — ООО «ИнженерИнвест»"), (5, "Заказчик")):
        ws.merge_cells(start_row=podpis, start_column=col,
                       end_row=podpis, end_column=col + 2 if col == 1 else 6)
        c = ws.cell(row=podpis, column=col)
        c.border = Border(top=Side(style="thin", color=INK))
        p = ws.cell(row=podpis + 1, column=col, value=text)
        p.font = Font(size=9, color=MUTED)

    # ---------- Печать ----------
    ws.print_area = f"A1:F{podpis + 1}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.freeze_panes = f"A{pervaya}"

    wb.calculation.fullCalcOnLoad = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    naydeno = dopisat_rezultaty(OUT, otvety)
    print(f"написано {OUT}: строк под позиции {vsego}, итог в F{ITOG}")
    print(f"ответы дописаны к {len(naydeno)} формулам")


if __name__ == "__main__":
    main()
