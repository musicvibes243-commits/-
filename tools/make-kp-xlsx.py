#!/usr/bin/env python3
"""Собирает docs/kp.xlsx — коммерческое предложение для Excel.

Пара к tools/make-smeta-xlsx.py. Смета — расчёт для разговора на объекте,
коммерческое предложение — официальная бумага с реквизитами, по которой
платят.

Суммы и итог — живые формулы. А вот сумма прописью формулой не считается:
в Excel для этого нужен макрос, а файл с макросом почта и мессенджеры
режут. Поэтому прописью пишется текстом при сборке файла.

**Если менять суммы в самом файле — прописью не пересчитается.**
Проще собрать новый: поправить список POZICII ниже и запустить скрипт.
Или заполнить kp.html — там прописью считается на лету.
Предупреждение об этом стоит в самом файле, за границей области печати:
на бумагу не попадёт, а при правке видно.

Запуск:  python3 tools/make-kp-xlsx.py
"""

import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from xlsx_formuly import dopisat_rezultaty

OUT = pathlib.Path("docs/kp.xlsx")

INK = "FF23303A"
MUTED = "FF7A8792"
ACCENT = "FFC97A09"
LINE = "FFD8D2C7"
LINE_2 = "FFECE7DE"

POZICII = [
    ("Котёл", 1, "шт", 49950),
    ("Монтаж", 1, "усл.", 49000),
    ("Расходный материал для котла", 1, "компл.", 13000),
    ("Отверстия", 2, "шт", 9500),
    ("Монтаж воздуховодов", 1, "усл.", 16000),
    ("Материал", 1, "компл.", 18000),
    ("Доставка материала", 1, "усл.", 7000),
]
ZAPAS = 5
DENGI = '#,##0.00" ₽"'

thin = Side(style="thin", color=LINE_2)
ramka = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---------- Сумма прописью ----------
ED = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять",
      "десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать",
      "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать"]
ED_J = ["", "одна", "две"]
DES = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят",
       "семьдесят", "восемьдесят", "девяносто"]
SOT = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот",
       "семьсот", "восемьсот", "девятьсот"]


def triada(n, jenskiy=False):
    out = []
    s, o = divmod(n, 100)
    if s:
        out.append(SOT[s])
    if o < 20:
        if o:
            out.append(ED_J[o] if jenskiy and o < 3 else ED[o])
    else:
        d, e = divmod(o, 10)
        out.append(DES[d])
        if e:
            out.append(ED_J[e] if jenskiy and e < 3 else ED[e])
    return " ".join(out)


def sklon(n, f1, f2, f5):
    o = n % 100
    if 10 < o < 20:
        return f5
    return {1: f1, 2: f2, 3: f2, 4: f2}.get(n % 10, f5)


def propisyu(summa):
    rub = int(summa)
    kop = round((summa - rub) * 100)
    if rub == 0:
        return f"Ноль рублей {kop:02d} копеек"
    gruppy = [(10 ** 9, False, ("миллиард", "миллиарда", "миллиардов")),
              (10 ** 6, False, ("миллион", "миллиона", "миллионов")),
              (10 ** 3, True, ("тысяча", "тысячи", "тысяч"))]
    out, ost = [], rub
    for d, j, f in gruppy:
        n, ost = divmod(ost, d)
        if n:
            out.append(triada(n, j))
            out.append(sklon(n, *f))
    if ost:
        out.append(triada(ost))
    t = " ".join(x for x in out if x)
    t = t[0].upper() + t[1:]
    return f"{t} {sklon(rub, 'рубль', 'рубля', 'рублей')} {kop:02d} копеек"


def rubli(n):
    """171950 → «171 950,00» — как это пишет Excel с русскими настройками."""
    return f"{n:,.2f}".replace(",", " ").replace(".", ",")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Предложение"

    # Готовые ответы формул: кладутся в файл рядом с самими формулами,
    # иначе быстрый просмотр на телефоне показывает нули. Подробности —
    # в tools/xlsx_formuly.py.
    otvety = {}

    for col, w in {"A": 5, "B": 44, "C": 9, "D": 10, "E": 14, "F": 16, "H": 60}.items():
        ws.column_dimensions[col].width = w

    def sliyanie(row, text, **kw):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(**kw) if kw else Font(color=INK)
        return c

    # ---------- Шапка ----------
    sliyanie(1, "СтройИнвест", bold=True, size=18, color=INK)
    sliyanie(2, "Инженерные системы дома и участка: септики, водоснабжение, "
                "отопление, электрика, дренаж", size=9, color=MUTED)
    sliyanie(3, "ООО «ИнженерИнвест» · ИНН 2100032923 · ОГРН 1262100002878",
             size=9, color=MUTED)
    sliyanie(4, "+7 (915) 346-97-28 · stroyinvest-mo.ru", size=9, color=MUTED)
    ws.row_dimensions[1].height = 24

    # ---------- Реквизиты ----------
    c = ws.cell(row=6, column=1, value="РЕКВИЗИТЫ ДЛЯ ОПЛАТЫ")
    c.font = Font(bold=True, size=9, color=ACCENT)
    rekv = [
        (7, "Банк", "", "БИК", ""),
        (8, "Кор. счёт", "", "Расч. счёт", ""),
        (9, "Получатель", "ООО «ИнженерИнвест»", "ИНН / ОГРН", "2100032923 / 1262100002878"),
    ]
    for row, p1, v1, p2, v2 in rekv:
        for col, podpis, znach in ((1, p1, v1), (4, p2, v2)):
            pc = ws.cell(row=row, column=col, value=podpis)
            pc.font = Font(size=8, color=MUTED)
            ws.merge_cells(start_row=row, start_column=col + 1,
                           end_row=row, end_column=col + 2)
            vc = ws.cell(row=row, column=col + 1, value=znach or None)
            vc.font = Font(size=10, color=INK)
            vc.border = Border(bottom=Side(style="dotted", color=LINE))

    # ---------- Заголовок документа ----------
    t = sliyanie(11, "КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ", bold=True, size=15, color=INK)
    t.alignment = Alignment(horizontal="center")
    ws.row_dimensions[11].height = 22

    for row, podpis, znach in ((12, "Номер", ""), (13, "Дата", ""),
                               (14, "Заказчик", "Частное лицо"),
                               (15, "Объект", ""),
                               (16, "Основание", "Запрос коммерческого предложения")):
        pc = ws.cell(row=row, column=1, value=podpis)
        pc.font = Font(size=8, color=MUTED)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        vc = ws.cell(row=row, column=2, value=znach or None)
        vc.font = Font(size=10, color=INK)
        vc.border = Border(bottom=Side(style="dotted", color=LINE))

    ws.merge_cells(start_row=18, start_column=1, end_row=19, end_column=6)
    isp = ws.cell(row=18, column=1,
                  value="Исполнитель: Общество с ограниченной ответственностью "
                        "«ИнженерИнвест», ИНН 2100032923, ОГРН 1262100002878, адрес: "
                        "429951, Чувашская Республика, г. Новочебоксарск, "
                        "ул. Комсомольская, д. 14, кв. 14. Работы выполняются "
                        "по Москве и Московской области.")
    isp.font = Font(size=9, color=INK)
    isp.alignment = Alignment(wrap_text=True, vertical="top")

    # ---------- Таблица ----------
    SHAPKA = 21
    for i, z in enumerate(["№", "Наименование", "Кол-во", "Ед.", "Цена, ₽", "Сумма, ₽"], 1):
        c = ws.cell(row=SHAPKA, column=i, value=z)
        c.font = Font(bold=True, size=9, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = ramka
    ws.row_dimensions[SHAPKA].height = 26

    pervaya = SHAPKA + 1
    vsego = len(POZICII) + ZAPAS
    for i in range(vsego):
        r = pervaya + i
        poz = POZICII[i] if i < len(POZICII) else ("", None, "", None)
        ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=poz[0] or None)
        ws.cell(row=r, column=3, value=poz[1]).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4, value=poz[2] or None).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5, value=poz[3]).number_format = DENGI
        s = ws.cell(row=r, column=6, value=f'=IF(OR(C{r}="",E{r}=""),"",C{r}*E{r})')
        s.number_format = DENGI
        # Пустая строка даёт пустой ответ, заполненная — произведение.
        otvety[f"F{r}"] = poz[1] * poz[3] if i < len(POZICII) else ""
        for col in range(1, 7):
            ws.cell(row=r, column=col).border = ramka
        ws.row_dimensions[r].height = 19

    poslednya = pervaya + vsego - 1

    ITOG = poslednya + 1
    ws.merge_cells(start_row=ITOG, start_column=1, end_row=ITOG, end_column=5)
    c = ws.cell(row=ITOG, column=1, value="ИТОГО")
    c.font = Font(bold=True, size=13, color=INK)
    c.alignment = Alignment(horizontal="right")
    c.border = Border(top=Side(style="medium", color=INK))
    otvety[f"F{ITOG}"] = sum(k * c for _, k, _, c in POZICII)
    it = ws.cell(row=ITOG, column=6, value=f"=SUM(F{pervaya}:F{poslednya})")
    it.font = Font(bold=True, size=13, color=INK)
    it.number_format = DENGI
    it.border = Border(top=Side(style="medium", color=INK))
    ws.row_dimensions[ITOG].height = 24

    # ---------- Прописью ----------
    itogo_summa = sum(k * c for _, k, _, c in POZICII)
    SLOV = ITOG + 2
    # Строка «Всего наименований…» — формула, а не число: иначе при правке цены
    # итог в таблице пересчитался бы, а эта строка осталась старой, и на одном
    # листе стояли бы две разные суммы.
    sliyanie(SLOV, f'="Всего наименований "&COUNTA(B{pervaya}:B{poslednya})'
                   f'&", на сумму "&TEXT(F{ITOG},"#,##0.00")&" ₽"',
             size=10, bold=True, color=INK)
    otvety[f"A{SLOV}"] = (f"Всего наименований {len(POZICII)}, "
                          f"на сумму {rubli(itogo_summa)} ₽")
    sl = sliyanie(SLOV + 1, propisyu(itogo_summa), size=10, color=INK)
    sl.font = Font(size=10, color=INK, underline="single")

    # Предупреждение — за границей области печати, на бумагу не попадёт.
    pred = ws.cell(row=SLOV + 1, column=8,
                   value="⚠ Сумма прописью не пересчитывается сама. Изменили цены — "
                         "перепишите строку вручную или соберите файл заново: "
                         "python3 tools/make-kp-xlsx.py. На странице kp.html "
                         "прописью считается на лету.")
    pred.font = Font(size=9, color="FFB00020")
    pred.alignment = Alignment(wrap_text=True, vertical="top")

    sliyanie(SLOV + 3, "Предложение действительно 14 дней с даты составления. "
                       "Окончательная стоимость и сроки фиксируются договором.",
             size=9, color=MUTED)

    podpis = SLOV + 6
    for col, text in ((1, "Руководитель ООО «ИнженерИнвест»"), (5, "Составил")):
        kon = 3 if col == 1 else 6
        ws.merge_cells(start_row=podpis, start_column=col, end_row=podpis, end_column=kon)
        ws.cell(row=podpis, column=col).border = Border(top=Side(style="thin", color=INK))
        p = ws.cell(row=podpis + 1, column=col, value=text)
        p.font = Font(size=9, color=MUTED)

    # ---------- Печать ----------
    ws.print_area = f"A1:F{podpis + 1}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)

    # Excel всё равно пересчитает формулы при открытии — записанные ответы
    # нужны только тем, кто файл не считает, а просто показывает.
    wb.calculation.fullCalcOnLoad = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    naydeno = dopisat_rezultaty(OUT, otvety)

    print(f"написано {OUT}: позиций {len(POZICII)}, итог {rubli(itogo_summa)} ₽")
    print(f"прописью: {propisyu(itogo_summa)}")
    print(f"ответы дописаны к {len(naydeno)} формулам")


if __name__ == "__main__":
    main()
